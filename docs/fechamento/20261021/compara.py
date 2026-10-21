"""Antes x depois do RH/DP: os VALORES têm de ser os mesmos (só muda apresentação). Não imprime nomes."""
import io, sys, datetime, collections
import openpyxl, decimal
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
A, D = openpyxl.load_workbook(sys.argv[1]), openpyxl.load_workbook(sys.argv[2])
REN = {"Comissao_Total": "Comissão Total", "Codigo_Operacao": "Código da Operação", "Plano_Financeiro": "Plano Financeiro", "Valor_SPF_Bruto": "Valor SPF Bruto",
       "Valor_SPF_70pct": "Valor SPF 70%", "Producao": "Produção", "SPF_Extra": "SPF Extra", "SPF_Liquido": "SPF Líquido", "Rentabilidade_Total": "Rentabilidade Total",
       "Comissao_Principal": "Comissão Principal", "Comissao_SPF": "Comissão SPF", "Observacao": "Observação", "Status": "Departamento"}
def norm(h, v, depois):
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    if h == "Share" and isinstance(v, (int, float)): return round(v * (100 if depois else 1), 6)
    if isinstance(v, (int, float)): return float(decimal.Decimal(repr(v)).quantize(decimal.Decimal("0.01"), rounding=decimal.ROUND_HALF_UP))
    return "" if v is None else v
ok_total = True
for wa, wd in zip(A.worksheets[1:], D.worksheets[1:]):
    ha = [REN.get(c.value, c.value) for c in wa[1]]; hd = [c.value for c in wd[1]]
    linhas = lambda ws, hs, dep: collections.Counter(tuple(norm(h, ws.cell(r, j + 1).value, dep) for j, h in enumerate(hs)) for r in range(2, ws.max_row + 1))
    ia, idd = linhas(wa, ha, False), linhas(wd, hd, True)
    ok = ha == hd and ia == idd; ok_total &= ok
    print(f"{wd.title:26} colunas iguais (após renomear): {ha == hd} | linhas: {sum(ia.values())} x {sum(idd.values())} | mesmos valores: {ia == idd}")
    if not ok and ha == hd:
        dif = (ia - idd) + (idd - ia); print("    diferenças:", sum(dif.values()), "| colunas afetadas:", sorted({hd[j] for t in list((ia - idd).keys())[:5] for j in range(len(hd)) if not any(t2[j] == t[j] for t2 in (idd - ia).keys())}))
r1 = {str(r[0].value): r[1].value for r in D.worksheets[0].iter_rows() if r[0].value and len(r) > 1}
somas = {wd.title: round(sum((wd.cell(r, [c.value for c in wd[1]].index("Comissão Total") + 1).value or 0) for r in range(2, wd.max_row + 1)), 2) for wd in D.worksheets[1:] if "Comissão Total" in [c.value for c in wd[1]]}
print("\nResumo Comissão Total:", r1.get("Comissão Total"), "| soma da coluna na Memória de Cálculo:", somas.get("8_MEMORIA_DE_CALCULO"), "| Vendedores+Analistas+Gerentes:", round(sum(v for k, v in somas.items() if k[0] in "234"), 2))
print("Fechado em:", r1.get("Fechado em"), "| Gerado em:", r1.get("Gerado em"))
print("\nRESULTADO:", "VALORES IGUAIS" if ok_total else "HÁ DIFERENÇA DE VALOR")
