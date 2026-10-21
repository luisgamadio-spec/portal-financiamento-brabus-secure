"""Inspeciona a estrutura do .xlsx do RH/DP sem imprimir nomes: cabeçalhos, formatos, congelamento, larguras, ordenação, totais."""
import io, sys, zipfile, re, collections
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
arq = sys.argv[1]
wb = openpyxl.load_workbook(arq)
z = zipfile.ZipFile(arq)
for i, ws in enumerate(wb.worksheets, 1):
    xml = z.read(f"xl/worksheets/sheet{i}.xml").decode("utf-8")
    pane = re.search(r"<pane[^>]*>", xml)
    print(f"\n=== {ws.title}  ({ws.max_row - 1} linhas)  congelado: {ws.freeze_panes or '-'}  {'(pane no XML)' if pane else '(sem pane no XML)'}")
    if ws.title.startswith("1_"):
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            a, b = row[0], (row[1] if len(row) > 1 else None)
            if a.value is None: continue
            print(f"   {str(a.value)[:34]:34} | {repr(b.value)[:30] if b is not None else '':30} | fmt={b.number_format if b is not None else ''} | tipo={b.data_type if b is not None else ''}")
        continue
    cab = [c.value for c in ws[1]]
    larg = [round(ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width or 0) for j in range(len(cab))]
    print("   colunas:", cab)
    fmts = []
    for j in range(len(cab)):
        tipos = collections.Counter((c.data_type, c.number_format) for c in [ws.cell(r, j + 1) for r in range(2, ws.max_row + 1)] if c.value not in (None, ""))
        fmts.append(tipos.most_common(1)[0][0] if tipos else ("vazio", ""))
    for h, f, w in zip(cab, fmts, larg): print(f"     {str(h)[:28]:28} tipo={f[0]} fmt={f[1]:22} largura={w}")
    idx = {h: j for j, h in enumerate(cab)}
    chave = [k for k in ("Loja", "Departamento", "Nome", "Vendedor") if k in idx]
    if chave and ws.max_row > 2:
        vals = [tuple(str(ws.cell(r, idx[k] + 1).value or "") for k in chave) for r in range(2, ws.max_row + 1)]
        import unicodedata
        norm = lambda x: "".join(ch for ch in unicodedata.normalize("NFD", x) if unicodedata.category(ch) != "Mn").casefold()
        esperado = sorted(vals, key=lambda t: tuple(norm(x) for x in t))
        fora = sum(1 for a, b in zip(vals, esperado) if tuple(map(norm, a)) != tuple(map(norm, b)))
        print("   ordenado por", chave, ":", fora == 0, f"({fora} posições fora)" if fora else "")
    for col in ("Share", "CPF"):
        if col in idx:
            v = [ws.cell(r, idx[col] + 1).value for r in range(2, ws.max_row + 1)]
            print(f"   {col}: preenchidos={sum(1 for x in v if x not in (None, ''))}", (f"máx={max(x for x in v if isinstance(x, (int, float)))}" if col == "Share" and any(isinstance(x, (int, float)) for x in v) else ""))
    for col in [c for c in cab if c and "Comiss" in str(c) and "Total" in str(c)]:
        print(f"   soma {col}:", round(sum((ws.cell(r, idx[col] + 1).value or 0) for r in range(2, ws.max_row + 1)), 2))
