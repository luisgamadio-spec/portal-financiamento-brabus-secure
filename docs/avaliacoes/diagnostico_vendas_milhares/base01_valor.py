import openpyxl, sys, io, collections, re, datetime
import openpyxl.reader.excel as _rx
_rx.apply_stylesheet = lambda archive, wb: None
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ws = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True).worksheets[0]
rows = ws.iter_rows(values_only=True)
head = None
for r in rows:
    if r and sum(isinstance(c, str) for c in r) >= 5: head = [str(c).strip() if c else "" for c in r]; break
iv, idt = head.index("Valor Venda"), next(i for i, h in enumerate(head) if h.lower().startswith("data venda"))
tipos = collections.Counter(); exemplos = collections.Counter(); n = 0; ate_maio = 0; texto_ponto = 0
for r in rows:
    if not r or r[iv] in (None, ""): continue
    n += 1; v = r[iv]; d = r[idt]
    if isinstance(d, str):
        try: d = datetime.datetime.strptime(d.strip()[:10], "%d/%m/%Y")
        except ValueError: d = None
    tipos[type(v).__name__] += 1
    if d and d < datetime.datetime(2026, 6, 1): ate_maio += 1
    if isinstance(v, str):
        exemplos[re.sub(r"\d", "9", v.strip())] += 1
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", v.strip()): texto_ponto += 1
print("linhas com valor:", n, "| vendas até maio:", ate_maio, "| tipos:", dict(tipos))
print("formatos em texto (dígitos trocados por 9):", exemplos.most_common(8))
print("textos só com ponto de milhar (ex. 179.990), que o importador lê errado:", texto_ponto)
ws = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True).worksheets[0]
rows = ws.iter_rows(values_only=True)
for r in rows:
    if r and sum(isinstance(c, str) for c in r) >= 5: break
peq = [r[iv] for r in rows if r and isinstance(r[iv], (int, float)) and r[iv] < 1000]
print("valores abaixo de 1.000 NO PRÓPRIO ARQUIVO:", len(peq), "| exemplos:", sorted(collections.Counter(peq).most_common(6)))
